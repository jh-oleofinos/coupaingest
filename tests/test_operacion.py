"""Pruebas de las reglas operativas: ocultas, vacías, candado, archivo sin cambios."""

import io

import openpyxl
import pytest

from coupa_ingest import pipeline as pl
from coupa_ingest.config import Settings
from coupa_ingest.transform import ValidationError, excel_to_csvs


def _book_with_hidden_and_empty() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Visible"
    ws.append(["id"])
    ws.append(["1"])
    oculta = wb.create_sheet("Oculta")
    oculta.append(["id"])
    oculta.append(["2"])
    oculta.sheet_state = "hidden"
    vacia = wb.create_sheet("Vacia")
    vacia.append(["id"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_hidden_skipped_and_empty_fails(tmp_path):
    xlsx = _book_with_hidden_and_empty()
    with pytest.raises(ValidationError, match="Vacia"):
        excel_to_csvs(xlsx, [], tmp_path, all_sheets=True)
    res = excel_to_csvs(xlsx, [], tmp_path, all_sheets=True, exclude=["Vacia"])
    assert [r.sheet for r in res] == ["Visible"]
    res = excel_to_csvs(xlsx, [], tmp_path, all_sheets=True, include_hidden=True, allow_empty=True)
    assert [r.sheet for r in res] == ["Visible", "Oculta", "Vacia"]


def _settings(tmp_path, **kw) -> Settings:
    return Settings(
        data_dir=tmp_path / "data",
        log_dir=tmp_path / "logs",
        sharepoint={"tenant_id": "t", "client_id": "c", "hostname": "h", "site_path": "/s", "file_path": "/f.xlsx"},
        coupa={"base_url": "https://coupa.invalid", "model_id": "M1"},
        **kw,
    )


def test_lock_blocks_second_run(tmp_path, sample_xlsx):
    s = _settings(tmp_path)
    (s.data_dir).mkdir(parents=True)
    (s.data_dir / ".lock").write_text("123")
    src = tmp_path / "in.xlsx"
    src.write_bytes(sample_xlsx)
    m = pl.run(s, "r1", dry_run=True, local_file=src)
    assert m.status == "failed" and "en curso" in m.error
    assert (s.data_dir / ".lock").exists()   # no borra el candado ajeno


def test_unchanged_file_skip(tmp_path, sample_xlsx, monkeypatch):
    class FakeSP:
        def __init__(self, *a, **k):
            pass

        def download_file(self):
            return sample_xlsx, {"name": "f.xlsx", "eTag": "v1"}

    monkeypatch.setattr(pl, "SharePointClient", FakeSP)
    monkeypatch.setattr(pl.secrets, "get_secret", lambda n, a: "x")
    s = _settings(tmp_path, if_unchanged="skip")
    s.data_dir.mkdir(parents=True)
    (s.data_dir / "state.json").write_text('{"run_id": "prev", "source": {"eTag": "v1"}}')
    m = pl.run(s, "r2")
    assert m.status == "ok" and m.skipped


def test_placeholders_detected(tmp_path):
    s = _settings(tmp_path)
    s.sharepoint.file_path = "/PENDIENTE/x.xlsx"
    s.sharepoint.tenant_id = "00000000-0000-0000-0000-000000000000"
    assert set(s.placeholders()) == {"sharepoint.file_path", "sharepoint.tenant_id"}
